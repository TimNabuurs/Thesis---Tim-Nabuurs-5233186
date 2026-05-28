import pandas as pd
import numpy as np
from TES_ZoneModel_Class import ZoneModel
from TES_Availability_Engine import calculate_thermal_potentials
from openpyxl.utils import get_column_letter

def run_dynamic_building_simulation(start_day=0, end_day=365, silent_mode=False, export_excel=False):
    """
    Executes the multi-zone thermodynamic building loop using a Rolling Horizon MPC.
    Dynamically generates and changes the UFH schedule block every single day
    for BOTH heating (winter) and cooling (summer) configurations based on 24-hour weather, 
    kVA forecasts, and high-fidelity ZoneModel predictive sweeps.
    """
    h_day = [21.0] * 24
    c_day = [21.0] * 24
    heating_profile = (h_day * 7) 
    cooling_profile = (c_day * 7) 

    h_atrium = [16.0] * 24
    c_atrium = [28.0] * 24
    atrium_h_profile = h_atrium * 7
    atrium_c_profile = c_atrium * 7

    base_inputs = {
        "t_ground": 15.0,
        "heating_setpoint_profile": heating_profile,
        "cooling_setpoint_profile": cooling_profile,
        "vent_comfort_limit": 21.0, 
        "initial_temp": 21.0,
        "heating_setpoint": 21.0,
        "cooling_setpoint": 21.0,
        "rc_roof": 7.84,
        "rc_ground_floor": 4.7,
        "u_value_windows": 1.84,
        "alfai": 7.5,
        "alfao": 15.0,
        "thermal_mass_factor": 450000,
        "solar_absorption_coefficient": 0.5,
        "solar_heat_coefficient_shading": 0.20,
        "solar_heat_coefficient_glazing": 0.35,
        "air_density": 1.2,
        "air_heat_capacity": 1003.0,
        "vent_flow_per_person": 36.0,
        "infiltration_ach": 0.2,
        "natural_vent_rate": 3.0,
        "heat_per_person": 65.0,
        "appliances_w_m2": 6.0,
        "lighting_w_m2": 6.0,
        "heating_power_max": 1000000.0,     
        "cooling_power_max": -1000000.0,    
        "system_pressure_drop": 200.0,
        "efficiency_fan_and_motor": 0.6,
        "eta": 0.0, 
        "u_floor_heating": 11.0,                
        "floor_heating_area": 200,              
        "t_floor_water_heating": 35.0,          
        "t_floor_water_cooling": 18.0,          
        "t_floor_limit_heating": 24, 
        "t_floor_limit_cooling": 18.0, 
        "system_profile": [1.0]*168
    }

    internal_wall_connections = [
        {"zones": ("Zone 1", "Zone 2"), "area": 505.0, "R": 0.33},  
        {"zones": ("Zone 1", "Zone 4"), "area": 122.0, "R": 0.60},  
        {"zones": ("Zone 1", "Zone 6"), "area": 35.0, "R": 1.00},   
        {"zones": ("Zone 3", "Zone 2"), "area": 355.0, "R": 0.33},  
        {"zones": ("Zone 3", "Zone 4"), "area": 295.0, "R": 0.33}, 
        {"zones": ("Zone 3", "Zone 5"), "area": 525.0, "R": 0.50},  
        {"zones": ("Zone 3", "Zone 6"), "area": 177.0, "R": 1.00}, 
        {"zones": ("Zone 4", "Zone 2"), "area": 208.0, "R": 0.50},
        {"zones": ("Zone 4", "Zone 5"), "area": 20.0, "R": 0.50},  
        {"zones": ("Zone 4", "Zone 6"), "area": 218.0, "R": 1.00}, 
        {"zones": ("Zone 5", "Zone 6"), "area": 266.0, "R": 1.00}, 
        {"zones": ("Zone 6", "Zone 2"), "area": 359.0, "R": 1.00}, 
    ]

    adjacencies = []
    for connection in internal_wall_connections:
        adjacencies.append({"zones": connection["zones"], "UA": connection["area"] / connection["R"]})

    zone_definitions = []
    
    # Zone 1
    z1 = base_inputs.copy()
    z1.update({"floor_area": 455.0, "area_roof": 455.0, "area_ground": 455.0, "room_volume": 4077.0, "rc_facade": 4.86, "max_people_per_m2": 0.55, "appliances_w_m2": 12.0, "lighting_w_m2": 15.0, "floor_heating_activated": False, "total_facade_areas": {'N': 255.0, 'NE': 0.0, 'E': 0.0, 'SE': 0.0, 'S': 0.0, 'SW': 0.0, 'W': 0.0, 'NW': 0.0}, "window_percentages": {'N': 0.0, 'NE': 0.0, 'E': 0.0, 'SE': 0.0, 'S': 0.0, 'SW': 0.0, 'W': 0.0, 'NW': 0.0}, "glazing_percentages": {'N': 0.0, 'NE': 0.0, 'E': 0.0, 'SE': 0.0, 'S': 0.0, 'SW': 0.0, 'W': 0.0, 'NW': 0.0}})
    z1["occ_profile"] = [0.0]*48 + [0.0]*15 + [0.4, 0.4, 0.4, 0.4, 1.0, 1.0, 1.0, 1.0, 0.4] + [0.0]*24 + [0.0]*15 + [0.4, 0.4, 0.4, 0.4, 1.0, 1.0, 1.0, 1.0, 0.4] + [0.0]*15 + [0.4, 0.4, 0.4, 0.4, 1.0, 1.0, 1.0, 1.0, 0.4] + [0.0]*15 + [0.4, 0.4, 0.4, 0.4, 1.0, 1.0, 1.0, 1.0, 0.4]
    z1["equip_profile"] = z1["occ_profile"]
    z1["vent_profile"] = ([0.10]*8 + [1.00]*16)*7
    zone_definitions.append(("Zone 1", z1))

    # Zone 2
    z2 = base_inputs.copy()
    z2.update({"floor_area": 565.0, "area_roof": 565.0, "area_ground": 565.0, "room_volume": 7370.0, "rc_roof": 13.66, "rc_facade": 4.7, "max_people_per_m2": 0.28, "appliances_w_m2": 10.0, "lighting_w_m2": 6.0, "floor_heating_activated": True, "heating_setpoint_profile": atrium_h_profile, "cooling_setpoint_profile": atrium_c_profile, "heating_setpoint": 16.0, "cooling_setpoint": 28.0, "vent_cooling_setpoint": 27.0, "total_facade_areas": {'N': 0.0, 'NE': 80.0, 'E': 0.0, 'SE': 205.0, 'S': 0.0, 'SW': 315.0, 'W': 0.0, 'NW': 230.0}, "window_percentages": {'N': 0.0, 'NE': 0.0, 'E': 0.00, 'SE': 0.293, 'S': 0.0, 'SW': 0.143, 'W': 0.0, 'NW': 0.187}, "glazing_percentages": {'N': 0.0, 'NE': 0.0, 'E': 0.0, 'SE': 0.80, 'S': 0.0, 'SW': 0.80, 'W': 0.0, 'NW': 0.80}})
    z2["occ_profile"] = ([0.0]*8 + [0.2]*10 + [0.3]*6)*2 + ([0.0]*8 + [0.2]*7 + [0.5]*4 + [0.3]*4 + [0.5]) + ([0.0]*8 + [0.2]*10 + [0.3]*6) + ([0.0]*8 + [0.2]*7 + [0.5]*4 + [0.3]*4 + [0.3]) + ([0.0]*8 + [0.2]*7 + [0.5]*4 + [0.3]*4 + [0.5])*2
    z2["equip_profile"] = z2["occ_profile"]
    z2["vent_profile"] = ([0.10]*8 + [1.00]*16)*7
    zone_definitions.append(("Zone 2", z2))

    # Zone 3
    z3 = base_inputs.copy()
    z3.update({"floor_area": 685.0, "area_roof": 160.0, "area_ground": 525.0, "room_volume": 4315.0, "rc_facade": 4.86, "max_people_per_m2": 0.26, "appliances_w_m2": 10.0, "lighting_w_m2": 12.0, "floor_heating_activated": False, "vent_flow_per_person": 25.2, "total_facade_areas": {'N': 0.0, 'NE': 93.0, 'E': 0.0, 'SE': 265.0, 'S': 0.0, 'SW': 0.0, 'W': 0.0, 'NW': 64.0}, "window_percentages": {'N': 0.0, 'NE': 0.032, 'E': 0.0, 'SE': 0.0, 'S': 0.0, 'SW': 0.0, 'W': 0.0, 'NW': 0.0}, "glazing_percentages": {'N': 0.00, 'NE': 0.80, 'E': 0.00, 'SE': 0.00, 'S': 0.00, 'SW': 0.00, 'W': 0.00, 'NW': 0.80}})
    z3["occ_profile"] = ([0.0]*12 + [0.4]*7 + [1.0]*4 + [0.0])*7
    z3["equip_profile"] = z3["occ_profile"]
    z3["vent_profile"] = ([0.10]*8 + [1.00]*16)*7
    zone_definitions.append(("Zone 3", z3))

    # Zone 4
    z4 = base_inputs.copy()
    z4.update({"floor_area": 730.0, "area_roof": 295.0, "area_ground": 218.0, "room_volume": 2407.0, "rc_facade": 4.86, "max_people_per_m2": 0.06, "appliances_w_m2": 7.5, "lighting_w_m2": 6.0, "floor_heating_activated": False, "vent_flow_per_person": 25.2, "total_facade_areas": {'N': 0.0, 'NE': 13.0, 'E': 0.0, 'SE': 80.0, 'S': 0.0, 'SW': 150.0, 'W': 0.0, 'NW': 95.0}, "window_percentages": {'N': 0.0, 'NE': 0.0, 'E': 0.0, 'SE': 0.475, 'S': 0.0, 'SW': 0.00, 'W': 0.00, 'NW': 0.642}, "glazing_percentages": {'N': 0.00, 'NE': 0.00, 'E': 0.00, 'SE': 0.80, 'S': 0.00, 'SW': 0.00, 'W': 0.00, 'NW': 0.80}})
    z4["occ_profile"] = ([0.0]*8 + [1.0]*4 + [0.5] + [1.0]*4 + [0.0]*7)*5 + [0.0]*24*2
    z4["equip_profile"] = z4["occ_profile"]
    z4["vent_profile"] = ([0.10]*8 + [1.00]*16)*7
    zone_definitions.append(("Zone 4", z4))

    # Zone 5
    z5 = base_inputs.copy()
    z5.update({"floor_area": 480.0, "area_roof": 319.0, "area_ground": 160.0, "room_volume": 2495.0, "rc_facade": 4.86, "max_people_per_m2": 0.09, "appliances_w_m2": 5.0, "lighting_w_m2": 6.0, "floor_heating_activated": False, "vent_flow_per_person": 25.2, "total_facade_areas": {'N': 0.0, 'NE': 176.0, 'E': 0.0, 'SE': 110.0, 'S': 0.0, 'SW': 0.0, 'W': 0.0, 'NW': 74.0}, "window_percentages": {'N': 0.0, 'NE': 0.034, 'E': 0.0, 'SE': 0.0, 'S': 0.0, 'SW': 0.0, 'W': 0.0, 'NW': 0.0}, "glazing_percentages": {'N': 0.00, 'NE': 0.80, 'E': 0.00, 'SE': 0.00, 'S': 0.00, 'SW': 0.00, 'W': 0.00, 'NW': 0.80}})
    z5["occ_profile"] = z4["occ_profile"]
    z5["equip_profile"] = z5["occ_profile"]
    z5["vent_profile"] = ([0.10]*8 + [1.00]*16)*7
    zone_definitions.append(("Zone 5", z5))

    # Zone 6
    z6 = base_inputs.copy()
    z6.update({"floor_area": 420.0, "area_roof": 360.0, "area_ground": 283.0, "room_volume": 3456.0, "rc_facade": 4.86, "max_people_per_m2": 0.03, "appliances_w_m2": 0.0, "lighting_w_m2": 0.0, "floor_heating_activated": False, "vent_flow_per_person": 25.2, "heating_setpoint_profile": [16.0]*168, "cooling_setpoint_profile": [28.0]*168, "vent_cooling_setpoint": 27.0, "total_facade_areas": {'N': 0.0, 'NE': 130.0, 'E': 0.0, 'SE': 133.0, 'S': 0.0, 'SW': 53.0, 'W': 0.0, 'NW': 75.0}, "window_percentages": {'N': 0.0, 'NE': 0.0, 'E': 0.0, 'SE': 0.083, 'S': 0.0, 'SW': 0.0, 'W': 0.0, 'NW': 0.0}, "glazing_percentages": {'N': 0.00, 'NE': 0.00, 'E': 0.00, 'SE': 0.80, 'S': 0.00, 'SW': 0.00, 'W': 0.00, 'NW': 0.00}})
    z6["occ_profile"] = ([0.0]*8 + [0.3]*14 + [0.1]*2)*2 + ([0.0]*13 + [0.7]*2 + [1.0]*9) + ([0.0]*8 + [0.3]*14 + [0.1]*2) + ([0.0]*13 + [0.7]*2 + [1.0]*9)*3
    z6["equip_profile"] = z6["occ_profile"]
    z6["vent_profile"] = ([0.10]*8 + [1.00]*16)*7
    z6["system_profile"] = [0.0]*168
    zone_definitions.append(("Zone 6", z6))

    file_path = 'Input Multi Zone Model - NEN5060-B2 1% solar orient - Extreme Year.csv'
    weather_df = pd.read_csv(file_path, sep=';')
    weather_df.columns = weather_df.columns.str.strip()
    if 'temp_ext' in weather_df.columns and 'T' not in weather_df.columns:
        weather_df = weather_df.rename(columns={'temp_ext': 'T'})
        
    df_solar_data = pd.read_csv('Input PV - NEN5060-B2 1%.txt', sep=',')
    weather_df['GG'] = df_solar_data['GG'].values[:len(weather_df)]
    weather_df['GD'] = df_solar_data['GD'].values[:len(weather_df)]
    
    TOTAL_GRID_LIMIT = 61.25
    grid_kva_history = []
    for t in range(8760):
        total_base_load_w = 0
        idx_168 = t % 168
        for name, params in zone_definitions:
            light_sch = params['equip_profile'][idx_168]
            vent_sch  = params['vent_profile'][idx_168]
            sys_active = params.get('system_profile', [1.0]*168)[idx_168]
            
            z_load = (params['lighting_w_m2'] + params['appliances_w_m2']) * params['floor_area'] * light_sch
            max_people = params['floor_area'] * params['max_people_per_m2']
            fans_power = (max_people * params['vent_flow_per_person'] * params['system_pressure_drop']) / (3600 * params['efficiency_fan_and_motor'])
            z_load += (fans_power * vent_sch) + ((0.5 * params['floor_area']) if sys_active > 0 else 0)
            total_base_load_w += z_load

        building_kva = total_base_load_w / 1000.0
        grid_kva_history.append(max(0, TOTAL_GRID_LIMIT - building_kva))

    total_h_baseline = np.zeros(len(weather_df))
    total_c_baseline = np.zeros(len(weather_df))

    for t_pass in range(len(weather_df)):
        is_winter_pass = weather_df['T'].iloc[t_pass:min(t_pass+24, len(weather_df))].mean() < 15.5
        row_temp = weather_df['T'].iloc[t_pass]
        
        for name, params in zone_definitions:
            if is_winter_pass:
                total_h_baseline[t_pass] += max(0, (21.0 - row_temp) * params['floor_area'] * 2.0 / 1000.0)
            else:
                total_c_baseline[t_pass] += max(0, (row_temp - 21.0) * params['floor_area'] * 2.0 / 1000.0)

    df_grid_all = pd.DataFrame({'Grid_kVA': grid_kva_history})
    global_potentials = calculate_thermal_potentials(
        weather_df.copy(), df_grid_all, split_mode='proportional',
        h_demand=pd.Series(total_h_baseline), c_demand=pd.Series(total_c_baseline)
    )
        
    HP_MAX_CAPACITY_H = 115.0
    HP_MAX_CAPACITY_C = 65.0
    
    # Extract both structural system limits cleanly for the dual-mode lookup bounds
    hp_ceilings_heating = np.minimum(global_potentials['Potential_Heating_kW'].values, HP_MAX_CAPACITY_H)
    hp_ceilings_cooling = np.minimum(global_potentials['Potential_Cooling_kW'].values, HP_MAX_CAPACITY_C)
    outdoor_temps_global = weather_df['T'].values

    if not silent_mode:
        print("Starting FULL DUAL-MODE Dynamic Daily Optimization Loop (365 Days)...")
        
    # --- NEW ---
    zones = {name: ZoneModel(name, **{**base_inputs, **params, 'hourly_df': weather_df}) 
            for name, params in zone_definitions}
    current_midnight_temps = {name: base_inputs['initial_temp'] for name in zones.keys()}
    # --- NEW ---

    final_dynamic_ufh_schedule = np.zeros(8760)
    cumulative_annual_deficit = 0

    for day in range(365):
        day_start_hour = day * 24
        day_end_hour = day_start_hour + 24
        
        day_temps = outdoor_temps_global[day_start_hour:day_end_hour]
        is_winter_day = day_temps.mean() < 15.5

        # --- NEW ---
        for name, zone_obj in zones.items():
            zone_obj.current_temp = current_midnight_temps[name]
        # --- NEW ---

        # Select appropriate matching tracking target ceilings depending on the active season
        day_ceilings = hp_ceilings_heating[day_start_hour:day_end_hour] if is_winter_day else hp_ceilings_cooling[day_start_hour:day_end_hour]
        
        best_daily_start = 0
        best_daily_duration = 0
        min_daily_deficit_score = float('inf')
        
        # --- HIGH SPEED ARCHITECTURE UPDATE: ELIMINATED DEEPCOPY
        midnight_temps = {name: zone_obj.current_temp for name, zone_obj in zones.items()}
        midnight_history_len = len(zones['Zone 1'].annual_heating_results)

        # --- REDUCED RANGE: Optimize from 0 hours up to 8 hours maximum
        for candidate_duration in range(0, 9):
            for candidate_start in range(24):
                day_bits = np.zeros(24)
                if candidate_duration > 0:
                    for i in range(candidate_duration):
                        day_bits[(candidate_start + i) % 24] = 1.0
                
                # DEBUG
                # if sum(day_bits) > 0:
                #   print(f"DEBUG: MPC testing Duration {candidate_duration} at Start {candidate_start}")
                # DEBUG
                
                # Fast memory float reset instead of copy.deepcopy
                for name, zone_obj in zones.items():
                    zone_obj.current_temp = midnight_temps[name]
                
                predicted_day_deficit = 0
                
                for hour_offset in range(24):
                    global_hour = day_start_hour + hour_offset

                    is_ufh_active = day_bits[hour_offset] > 0
                    
                    flows = {name: 0.0 for name in zones.keys()}
                    for adj in adjacencies:
                        z_a, z_b = adj["zones"]
                        f = adj["UA"] * (zones[z_b].current_temp - zones[z_a].current_temp)
                        flows[z_a] += f
                        flows[z_b] -= f
                        
                    atrium_temp = zones['Zone 2'].current_temp
                    total_hour_load_kw = 0
                    
                    for name, zone_obj in zones.items():
                        if name == 'Zone 2':
                            
                            # DEBUG
                            # print(f"DEBUG: Setting {name} (ID: {id(zone_obj)}) to {is_ufh_active}")
                            # DEBUG

                            zone_obj.floor_heating_activated = is_ufh_active
                        
                        # DEBUG
                        # is_active = bool(final_dynamic_ufh_schedule[t] != 0)
                        # zone_obj.floor_heating_activated = is_active
                        # if is_active:
                        #    print(f"DEBUG: Orchestrator set UFH to True for Zone 2 at hour {t}")
                        # DEBUG
                        
                        source = None if name == 'Zone 2' else atrium_temp
                        zone_obj.calculate_hour_step(global_hour, external_q_flow=flows[name], t_source=source, is_winter=is_winter_day)
                        
                        # NEW
                        # Accumulate correct matching operational metrics for cost mapping
                        if is_winter_day:
                            load = zone_obj.annual_heating_results[-1]
                        else:
                            load = zone_obj.annual_cooling_results[-1]

                        if name == 'Zone 2' and zone_obj.floor_heating_activated:
                            ufh_energy = zone_obj.detailed_results.get('UFH_kWh', [0.0])[-1]        # BEFORE IT WAS: [0.0] * (t+ 1))[t]
                            load += ufh_energy

                        total_hour_load_kw += load
                        # NEW

                    if is_winter_day:
                        # Heating deficit: Demand exceeds ceiling
                        hourly_deficit = max(0.0, total_hour_load_kw - day_ceilings[hour_offset])
                    else:
                        # Cooling deficit: Cooling demand magnitude exceeds ceiling magnitude
                        # We compare absolute values because cooling is negative
                        hourly_deficit = max(0.0, abs(total_hour_load_kw) - abs(day_ceilings[hour_offset]))
                    
                    predicted_day_deficit += (hourly_deficit ** 2)
                
                if predicted_day_deficit < min_daily_deficit_score:
                    min_daily_deficit_score = predicted_day_deficit
                    best_daily_start = candidate_start
                    best_daily_duration = candidate_duration
                
                # Trim out layout temporary simulation trial rows
                for name, zone_obj in zones.items():
                    zone_obj.annual_heating_results = zone_obj.annual_heating_results[:midnight_history_len]
                    zone_obj.annual_cooling_results = zone_obj.annual_cooling_results[:midnight_history_len]
                    zone_obj.annual_temp_results = zone_obj.annual_temp_results[:midnight_history_len]
                    for k in zone_obj.detailed_results.keys():
                        zone_obj.detailed_results[k] = zone_obj.detailed_results[k][:midnight_history_len]
        
        if best_daily_duration > 0:
            for i in range(best_daily_duration):
                final_dynamic_ufh_schedule[day_start_hour + ((best_daily_start + i) % 24)] = 1.0

        # Execute true pass for the day
        for hour_offset in range(24):
            t = day_start_hour + hour_offset
            flows = {name: 0.0 for name in zones.keys()}
            for adj in adjacencies:
                z_a, z_b = adj["zones"]
                f = adj["UA"] * (zones[z_b].current_temp - zones[z_a].current_temp)
                flows[z_a] += f
                flows[z_b] -= f
                
            atrium_temp = zones['Zone 2'].current_temp
            total_hour_load_kw = 0
            
            for name, zone_obj in zones.items():
                if name == 'Zone 2':
                    zone_obj.floor_heating_activated = bool(final_dynamic_ufh_schedule[t] != 0)
                
                source = None if name == 'Zone 2' else atrium_temp
                zone_obj.calculate_hour_step(t, external_q_flow=flows[name], t_source=source, is_winter=is_winter_day)
                
                # NEW
                if is_winter_day:
                    load = zone_obj.annual_heating_results[-1]
                else:
                    load = zone_obj.annual_cooling_results[-1]

                # Changed global hour to t
                if name == 'Zone 2' and zone_obj.floor_heating_activated:
                    ufh_energy = zone_obj.detailed_results.get('UFH_kWh', [0.0])[-1]        # BEFORE IT WAS: [0.0] * (t+ 1))[t]
                    load += ufh_energy

                total_hour_load_kw += load
                # NEW
                
            current_ceiling = hp_ceilings_heating[t] if is_winter_day else hp_ceilings_cooling[t]
            if total_hour_load_kw > current_ceiling:
                cumulative_annual_deficit += (total_hour_load_kw - current_ceiling)

        # --- NEW ---
        current_midnight_temps = {name: zone_obj.current_temp for name, zone_obj in zones.items()}
        # --- NEW ---

        # if not silent_mode and (day + 1 == 1 or (day + 1) % 15 == 0 or day + 1 == 365):
        #    mode_str = "HEATING" if is_winter_day else "COOLING"
        #   print(f" -> Day {day+1:03d}/365 | Mode: {mode_str} | Start: {best_daily_start:02d}:00 | Duration: {best_daily_duration:02d} hrs | Cum. Deficit: {cumulative_annual_deficit:.2f} kWh")

        if not silent_mode:
            mode_str = "HEATING" if is_winter_day else "COOLING"
            print(f" -> Day {day+1:03d}/365 | Mode: {mode_str} | Start: {best_daily_start:02d}:00 | Duration: {best_daily_duration:02d} hrs | Cum. Deficit: {cumulative_annual_deficit:.2f} kWh")

    # --- NEW ---
    hourly_export_data = []
    total_captured = len(list(zones.values())[0].annual_heating_results)
    for t in range(total_captured):
        row = {"Hour": t + 1}
        for name, zone_obj in zones.items():
            if name == "Zone 6": continue
            
            # 1. HVAC Demand
            heating = zone_obj.annual_heating_results[t]
            cooling = zone_obj.annual_cooling_results[t]
            net_demand = heating + cooling

            # 2. Update Demand if it is Zone 2 
            if name == "Zone 2":
                ufh_usage = zone_obj.detailed_results.get('UFH_kWh', [0.0])[-1]        # BEFORE IT WAS: [0.0] * (t+ 1))[t]
                net_demand += ufh_usage
                row["UFH_Energy_Zone 2"] = round(ufh_usage, 3)

            # 3. Write the total to the row
            row[f"Demand_{name}"] = round(net_demand, 1)
                          
        hourly_export_data.append(row)

    df_demands_combined = pd.DataFrame(hourly_export_data)
    df_demands_combined.to_csv("Output - Results_Yearly_Zone_Demands_Combined - DYNAMIC.csv", index=False)
        
    if not silent_mode:
        print("\n--- DYNAMIC OPTIMIZATION COMPLETE ---")
        print(f"Final Optimized Annual Tank Deficit: {cumulative_annual_deficit:.2f} kWh")
        print("Optimized dataset exported successfully to 'Output - Results_Yearly_Zone_Demands_Combined - DYNAMIC.csv'")

    if export_excel:
        export_filename = 'Building_Thermal_Performance_Detailed_DYNAMIC.xlsx'
        print(f"\nExporting detailed results to {export_filename}...")
        print("Please wait, calculating column widths for all zones...")

        final_results = []
        for name, zone_obj in zones.items():
            final_results.append({
                "Zone": name,
                "Annual Heating [kWh]": round(sum(zone_obj.annual_heating_results), 2),
                "Annual Cooling [kWh]": round(sum(zone_obj.annual_cooling_results), 2)
            })
        results_df = pd.DataFrame(final_results)

        with pd.ExcelWriter(export_filename, engine='openpyxl') as writer:
            results_df.to_excel(writer, sheet_name='Annual_Summary', index=False)
            
            summary_ws = writer.sheets['Annual_Summary']
            for column_cells in summary_ws.columns:
                content_length = max(len(str(cell.value)) for cell in column_cells)
                col_letter = get_column_letter(column_cells[0].column)
                summary_ws.column_dimensions[col_letter].width = min(max(content_length, 10), 25)
            
            for name, zone_obj in zones.items():
                zone_detailed_df = zone_obj.get_hourly_dataframe()
                zone_detailed_df['Q_cool (kWh)'] = zone_obj.annual_cooling_results
                zone_detailed_df.to_excel(writer, sheet_name=name, index=False)
                
                worksheet = writer.sheets[name]
                for column_cells in worksheet.columns:
                    content_length = max(len(str(cell.value)) for cell in column_cells)
                    col_letter = get_column_letter(column_cells[0].column)
                    worksheet.column_dimensions[col_letter].width = min(max(content_length, 10), 25)

        print(f"Export Complete. Detailed results saved in: {export_filename}")

    return df_demands_combined